import csv
import json
from pathlib import Path
from typing import List, Optional
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

from .models import ExtractionResult, CommentItem, ReactionItem
from .config import OUTPUT_DIR


def export_to_json(result: ExtractionResult, filepath: Optional[Path] = None) -> Path:
    """Export extraction result to a JSON file."""
    if not filepath:
        filepath = OUTPUT_DIR / f"{result.post.entity_type}_{result.post.entity_id}.json"
    
    data = result.model_dump()
    with open(filepath, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)
    
    return filepath


def export_to_csv(result: ExtractionResult, prefix: Optional[Path] = None) -> List[Path]:
    """Export comments and reactions into separate CSV files."""
    base_name = f"{result.post.entity_type}_{result.post.entity_id}"
    out_files = []

    # Export comments
    if result.comments:
        comments_file = OUTPUT_DIR / f"{base_name}_comments.csv"
        with open(comments_file, "w", encoding="utf-8-sig", newline="") as f:
            writer = csv.writer(f)
            writer.writerow([
                "Comment ID",
                "Author Name",
                "Headline",
                "Profile URL",
                "Comment Text",
                "Likes Count",
                "Replies Count",
                "Is Reply",
                "Parent Comment ID",
                "Created Desc",
                "Post URN",
            ])
            for c in result.comments:
                writer.writerow([
                    c.comment_id,
                    c.author.name,
                    c.author.headline,
                    c.author.profile_url,
                    c.text,
                    c.likes_count,
                    c.replies_count,
                    c.is_reply,
                    c.parent_comment_id or "",
                    c.created_at_desc or "",
                    c.post_urn,
                ])
        out_files.append(comments_file)

    # Export reactions
    if result.reactions:
        reactions_file = OUTPUT_DIR / f"{base_name}_reactions.csv"
        with open(reactions_file, "w", encoding="utf-8-sig", newline="") as f:
            writer = csv.writer(f)
            writer.writerow([
                "Reactor Name",
                "Headline",
                "Profile URL",
                "Reaction Type",
                "Post URN",
            ])
            for r in result.reactions:
                writer.writerow([
                    r.reactor.name,
                    r.reactor.headline,
                    r.reactor.profile_url,
                    r.reaction_type,
                    r.post_urn,
                ])
        out_files.append(reactions_file)

    return out_files


def export_to_excel(result: ExtractionResult, filepath: Optional[Path] = None) -> Path:
    """Export comments and reactions to Excel with styled headers."""
    if not filepath:
        filepath = OUTPUT_DIR / f"{result.post.entity_type}_{result.post.entity_id}.xlsx"

    wb = openpyxl.Workbook()
    # Remove default sheet
    wb.remove(wb.active)

    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="0077B5", end_color="0077B5", fill_type="solid") # LinkedIn Blue
    cell_alignment = Alignment(vertical="center")

    # 1. Comments Sheet
    ws_comments = wb.create_sheet(title="Comments")
    comment_headers = [
        "Author Name", "Headline", "Profile URL", "Comment Text",
        "Likes", "Replies", "Is Reply", "Created Time", "Comment ID", "Post URN"
    ]
    ws_comments.append(comment_headers)
    for col_num, _ in enumerate(comment_headers, 1):
        cell = ws_comments.cell(row=1, column=col_num)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = cell_alignment

    for c in result.comments:
        ws_comments.append([
            c.author.name,
            c.author.headline,
            c.author.profile_url,
            c.text,
            c.likes_count,
            c.replies_count,
            "Yes" if c.is_reply else "No",
            c.created_at_desc,
            c.comment_id,
            c.post_urn,
        ])

    # Auto column width
    for ws in [ws_comments]:
        for col in ws.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 60)

    # 2. Reactions Sheet
    ws_reactions = wb.create_sheet(title="Reactions")
    reaction_headers = ["Reactor Name", "Headline", "Profile URL", "Reaction Type", "Post URN"]
    ws_reactions.append(reaction_headers)
    for col_num, _ in enumerate(reaction_headers, 1):
        cell = ws_reactions.cell(row=1, column=col_num)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = cell_alignment

    for r in result.reactions:
        ws_reactions.append([
            r.reactor.name,
            r.reactor.headline,
            r.reactor.profile_url,
            r.reaction_type,
            r.post_urn,
        ])

    for col in ws_reactions.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws_reactions.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 60)

    wb.save(filepath)
    return filepath
